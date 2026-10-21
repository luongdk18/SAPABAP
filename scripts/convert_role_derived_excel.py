#!/usr/bin/env python3
"""
Convert Derived Roles Creation Excel file to flat row format for SAP Role upload.
Reads tabular sheet containing Master Roles, Derived Roles, and Descriptions.
Validates non-empty role keys and exports normalized CSV for SAP batch creation.

Output columns:
1. Master Role
2. Derived Role
3. Derived Role Description

Output:
- <filename>_converted.csv (UTF-8 with BOM for SAP GUI import)

Usage:
    python scripts/convert_role_derived_excel.py [input_excel_file]

Example:
    python scripts/convert_role_derived_excel.py excel-working/FICO_Roles_Derived.xlsx
    -> Outputs: excel-working/FICO_Roles_Derived_converted.csv
"""

import os
import sys
import csv
import openpyxl

# --- Safe Console Output on Windows ---
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')


def convert_role_derived_excel(input_path):
    """
    Reads Derived Roles Excel file and converts to flat CSV format.
    Dynamically maps header columns for Master Role, Derived Role, and Description.
    """
    # --- Path & File Resolution ---
    if not os.path.splitext(input_path)[1]:
        input_path = input_path + ".xlsx"

    if not os.path.exists(input_path):
        print(f"[ERROR] Input file not found: {input_path}")
        sys.exit(1)

    dir_name = os.path.dirname(input_path)
    base_name = os.path.basename(input_path)
    stem, _ = os.path.splitext(base_name)

    out_csv_name = f"{stem}_converted.csv"
    out_csv_path = os.path.join(dir_name, out_csv_name) if dir_name else out_csv_name

    # --- Load Workbook ---
    try:
        wb_in = openpyxl.load_workbook(input_path, data_only=True)
    except Exception as e:
        print(f"[ERROR] Could not read Excel workbook '{input_path}': {e}")
        sys.exit(1)

    # Locate 'DERIVED ROLES' sheet or fallback to active sheet
    target_sheet = None
    for name in wb_in.sheetnames:
        if "derived roles" in name.strip().lower():
            target_sheet = wb_in[name]
            break

    if target_sheet is None:
        target_sheet = wb_in.active
        print(f"[INFO] Using active sheet '{target_sheet.title}'.")
    else:
        print(f"[INFO] Processing sheet: '{target_sheet.title}'")

    ws = target_sheet

    # --- Dynamic Column Detection ---
    col_master = None
    col_derived = None
    col_desc = None

    for col in range(1, ws.max_column + 1):
        val = str(ws.cell(1, col).value or '').strip().lower()
        if 'master' in val:
            col_master = col
        elif 'derived' in val and ('desc' in val or 'text' in val or 'title' in val):
            col_desc = col
        elif 'derived' in val:
            col_derived = col

    # Fallback to standard positions (1: Master, 2: Derived, 3: Description)
    col_master = col_master or 1
    col_derived = col_derived or 2
    col_desc = col_desc or 3

    print(f"[INFO] Detected Columns -> Master Role: Col {col_master}, Derived Role: Col {col_derived}, Description: Col {col_desc}")

    # --- Row Extraction & Validation ---
    records = []
    skipped_empty = 0
    unique_masters = set()
    unique_derived = set()

    for row in range(2, ws.max_row + 1):
        raw_master = ws.cell(row, col_master).value
        raw_derived = ws.cell(row, col_derived).value
        raw_desc = ws.cell(row, col_desc).value

        master_role = str(raw_master or '').strip().upper()
        derived_role = str(raw_derived or '').strip().upper()
        role_desc = str(raw_desc) if raw_desc is not None else ''

        # Skip completely empty rows
        if not master_role and not derived_role and not role_desc.strip():
            skipped_empty += 1
            continue

        if not master_role or not derived_role:
            print(f"[WARNING] Row {row} missing Master or Derived Role: Master='{master_role}', Derived='{derived_role}' -> Skipped.")
            continue

        records.append({
            'master_role': master_role,
            'derived_role': derived_role,
            'role_desc': role_desc
        })
        unique_masters.add(master_role)
        unique_derived.add(derived_role)

    # --- CSV Export (UTF-8 with BOM) ---
    try:
        with open(out_csv_path, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            writer.writerow(["Master Role", "Derived Role", "Derived Role Description"])
            for item in records:
                writer.writerow([item['master_role'], item['derived_role'], item['role_desc']])
    except PermissionError:
        print(f"[ERROR] Permission denied when writing to '{out_csv_path}'. Please close the file in Excel and try again.")
        sys.exit(1)

    # --- Execution Summary ---
    print(f"\n--- Execution Summary ---")
    print(f"[OK] Input File       : {input_path}")
    print(f"[OK] Output CSV       : {out_csv_path} ({len(records)} data rows)")
    print(f"[OK] Master Roles     : {len(unique_masters)} unique")
    print(f"[OK] Derived Roles    : {len(unique_derived)} unique")
    if skipped_empty > 0:
        print(f"[INFO] Empty Rows     : {skipped_empty} empty row(s) skipped")


# --- CLI Entrypoint ---
if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] in ('-h', '--help'):
        print("Usage: python scripts/convert_role_derived_excel.py [path_to_excel_file]")
        print("Example: python scripts/convert_role_derived_excel.py excel-working/FICO_Roles_Derived.xlsx")
        sys.exit(0)

    target_file = None

    if len(sys.argv) > 1:
        target_file = sys.argv[1]
    else:
        # Default candidates lookup
        candidates = [
            os.path.join("excel-working", "FICO_Roles_Derived.xlsx"),
            os.path.join("excel-templates", "Template_Roles_Derived.xlsx"),
            "FICO_Roles_Derived.xlsx",
            "Roles_Derived.xlsx"
        ]
        for cand in candidates:
            if os.path.exists(cand):
                target_file = cand
                break

    if not target_file:
        print("[ERROR] No input file provided or found in default locations.")
        print("Usage: python scripts/convert_role_derived_excel.py <path_to_excel_file>")
        print("Example: python scripts/convert_role_derived_excel.py excel-working/FICO_Roles_Derived.xlsx")
        sys.exit(1)

    convert_role_derived_excel(target_file)
