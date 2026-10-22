#!/usr/bin/env python3
"""
Convert Derived Roles Matrix Excel file to flat row format for SAP upload.
Output format matches the CSV requirements for the SAP ABAP Derived Role Creation program.

Columns:
1. Master Role
2. Derived Role
3. Derived Role Description

Output:
- <filename>_converted.csv (UTF-8 with BOM for SAP GUI import)

Usage:
    python scripts/convert_role_02_derive.py [input_excel_file]

Example:
    python scripts/convert_role_02_derive.py excel-working/MM_Role_02_Derived.xlsx
    -> Outputs: excel-working/converted/MM_Role_02_Derived_converted.csv
"""

import os
import sys
import csv
import glob
import re
import openpyxl

# --- Safe Console Output on Windows ---
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')


def convert_role_02_derive(input_path, output_dir=None):
    """
    Reads Derived Roles Excel file and converts to flat CSV format.
    Expects columns:
    - Master / Leading Role (contains 'master', 'leading', or 'parent')
    - Derived Role (contains 'derive')
    - Derived Role Description (contains 'desc' or 'description')
    Output is saved to <parent_dir>/converted/ by default.
    """
    # --- Path & File Resolution ---
    if not os.path.splitext(input_path)[1]:
        input_path = input_path + ".xlsx"

    if not os.path.exists(input_path):
        print(f"[ERROR] Input file not found: {input_path}")
        return False

    base_name = os.path.basename(input_path)
    stem, _ = os.path.splitext(base_name)

    if not stem.endswith("_02_Derived"):
        out_csv_name = f"{stem}_02_Derived_converted.csv"
    else:
        out_csv_name = f"{stem}_converted.csv"

    if output_dir is None:
        dir_name = os.path.dirname(input_path) or "."
        output_dir = os.path.join(dir_name, "converted")

    if output_dir.endswith('.csv'):
        out_csv_path = output_dir
        if os.path.dirname(out_csv_path):
            os.makedirs(os.path.dirname(out_csv_path), exist_ok=True)
    else:
        os.makedirs(output_dir, exist_ok=True)
        out_csv_path = os.path.join(output_dir, out_csv_name)

    # --- Load Workbook ---
    try:
        wb_in = openpyxl.load_workbook(input_path, data_only=True)
    except Exception as e:
        print(f"[ERROR] Could not read Excel workbook '{input_path}': {e}")
        return False

    # Locate target sheet:
    # If consolidated role file (e.g. FICO_Role.xlsx, Template_Role.xlsx), Sheet 2 (Index 1) is Derived Role.
    # If standalone legacy file (e.g. MM_Role_02_Derived.xlsx), find sheet matching 'derive' or use active/Sheet 1.
    is_consolidated = not any(k in stem.lower() for k in ('_01_', '_02_', '_03_', '_04_', '_master', '_derive', '_org', '_object', '_composite'))

    if is_consolidated and len(wb_in.sheetnames) >= 2:
        target_sheet = wb_in.worksheets[1]
        print(f"[INFO] Consolidated file: Using Sheet 2 (Index 1): '{target_sheet.title}' for Derived Role data.")
    else:
        target_sheet = None
        for name in wb_in.sheetnames:
            if 'derive' in name.strip().lower():
                target_sheet = wb_in[name]
                break
        if target_sheet is None:
            target_sheet = wb_in.worksheets[0] if wb_in.worksheets else wb_in.active
            print(f"[INFO] Using sheet: '{target_sheet.title}'.")
        else:
            print(f"[INFO] Processing sheet: '{target_sheet.title}'")

    ws = target_sheet

    # --- Detect Header Columns ---
    header_row = 1
    col_master = None
    col_derived = None
    col_desc = None

    for col in range(1, ws.max_column + 1):
        val = str(ws.cell(header_row, col).value or '').strip().lower()
        if ('master' in val or 'leading' in val or 'parent' in val) and col_master is None:
            col_master = col
        elif 'derive' in val and ('desc' in val or 'description' in val) and col_desc is None:
            col_desc = col
        elif 'derive' in val and col_derived is None:
            col_derived = col

    # Fallback default positions: Col 1 = Master, Col 2 = Derived, Col 3 = Description
    if col_master is None:
        col_master = 1
    if col_derived is None:
        col_derived = 2
    if col_desc is None:
        col_desc = 3

    print(f"[INFO] Column mappings -> Master: Col {col_master}, Derived: Col {col_derived}, Desc: Col {col_desc}")

    records = []
    unique_masters = set()
    unique_derived = set()
    skipped_empty = 0

    # --- Scan Rows ---
    for row in range(header_row + 1, ws.max_row + 1):
        master_val = ws.cell(row, col_master).value
        derived_val = ws.cell(row, col_derived).value
        desc_val = ws.cell(row, col_desc).value

        master_str = str(master_val).strip() if master_val is not None else ''
        derived_str = str(derived_val).strip() if derived_val is not None else ''
        desc_str = str(desc_val).strip() if desc_val is not None else ''

        # Skip rows where both master and derived are empty
        if not master_str and not derived_str:
            skipped_empty += 1
            continue

        # If derived role exists, record it
        if derived_str:
            records.append({
                'master_role': master_str,
                'derived_role': derived_str,
                'derived_desc': desc_str
            })
            if master_str:
                unique_masters.add(master_str)
            unique_derived.add(derived_str)

    # --- CSV Export (UTF-8 with BOM) ---
    try:
        with open(out_csv_path, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            writer.writerow(["Master Role", "Derived Role", "Derived Role Description"])
            for item in records:
                writer.writerow([item['master_role'], item['derived_role'], item['derived_desc']])
    except PermissionError:
        print(f"[ERROR] Permission denied when writing to '{out_csv_path}'. Please close the file in Excel and try again.")
        return False

    # --- Execution Summary ---
    print(f"\n--- Execution Summary ---")
    print(f"[OK] Input File       : {input_path}")
    print(f"[OK] Output CSV       : {out_csv_path} ({len(records)} data rows)")
    print(f"[OK] Master Roles     : {len(unique_masters)} unique")
    print(f"[OK] Derived Roles    : {len(unique_derived)} unique")
    if skipped_empty > 0:
        print(f"[INFO] Empty Rows     : {skipped_empty} empty row(s) skipped")
    return True


# Compatibility alias
convert_role_derived_excel = convert_role_02_derive


# --- CLI Entrypoint ---
if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] in ('-h', '--help'):
        print("Usage: python scripts/convert_role_02_derive.py [path_to_excel_file]")
        print("Example: python scripts/convert_role_02_derive.py excel-working/MM_Role_02_Derived.xlsx")
        sys.exit(0)

    target_files = []

    if len(sys.argv) > 1:
        for arg in sys.argv[1:]:
            matches = [m for m in glob.glob(arg) if not os.path.basename(m).startswith('~$')]
            if matches:
                target_files.extend(matches)
            elif not os.path.basename(arg).startswith('~$'):
                target_files.append(arg)
    else:
        # Default candidates lookup in working and template directories
        # Prioritize consolidated *Role*.xlsx, fallback to legacy *_02_Derived.xlsx
        working_consolidated = [f for f in glob.glob(os.path.join("excel-working", "*Role*.xlsx"))
                                if not os.path.basename(f).startswith('~$')
                                and not any(k in os.path.basename(f) for k in ('_01_', '_02_', '_03_', '_04_', '_converted'))]
        working_legacy = [f for f in glob.glob(os.path.join("excel-working", "*_02_Derived*.xlsx"))
                          if not os.path.basename(f).startswith('~$')]
        
        consolidated_stems = {re.sub(r'_?Role.*$', '', os.path.basename(f), flags=re.I).upper() for f in working_consolidated}
        filtered_legacy = [f for f in working_legacy if re.sub(r'_?Role.*$', '', os.path.basename(f), flags=re.I).upper() not in consolidated_stems]
        
        unique_working = sorted(working_consolidated + filtered_legacy)
        
        if unique_working:
            target_files = unique_working
        else:
            tpl_consolidated = [f for f in glob.glob(os.path.join("excel-templates", "*Role*.xlsx"))
                                if not os.path.basename(f).startswith('~$')
                                and not any(k in os.path.basename(f) for k in ('_01_', '_02_', '_03_', '_04_', '_converted'))]
            tpl_legacy = [f for f in glob.glob(os.path.join("excel-templates", "*_02_Derived*.xlsx"))
                          if not os.path.basename(f).startswith('~$')]
            tpl_stems = {re.sub(r'_?Role.*$', '', os.path.basename(f), flags=re.I).upper() for f in tpl_consolidated}
            filtered_tpl = [f for f in tpl_legacy if re.sub(r'_?Role.*$', '', os.path.basename(f), flags=re.I).upper() not in tpl_stems]
            
            unique_templates = sorted(tpl_consolidated + filtered_tpl)
            if unique_templates:
                target_files = unique_templates
            else:
                target_files = ["Template_Role_02_Derived.xlsx"]

    if not target_files:
        print("[ERROR] No input file provided or found in default locations.")
        print("Usage: python scripts/convert_role_02_derive.py <path_to_excel_file>")
        print("Example: python scripts/convert_role_02_derive.py excel-working/MM_Role_02_Derived.xlsx")
        sys.exit(1)

    success = True
    for tf in target_files:
        print(f"\n==================================================")
        print(f"[RUNNING] Converting Derived Role: {tf}")
        print(f"==================================================")
        if not convert_role_02_derive(tf):
            success = False

    if not success:
        sys.exit(1)
