#!/usr/bin/env python3
"""
Convert User Role Assignment Matrix Excel file to flat row format for SAP upload.
Filters user-role assignments where cell value is marked with an indicator ('Y', 'y', 'X', 'x').
Generates normalized flat CSV format: Username, Role.
Prepares input file for ABAP batch role assignment program (via BAPI_USER_ACTGROUPS_ASSIGN / AGR_USERS).

Output columns:
1. Username
2. Role

Output:
- <filename>_converted.csv (UTF-8 with BOM for SAP GUI import)

Usage:
    python scripts/convert_users_assign_roles_excel.py [input_excel_file]

Example:
    python scripts/convert_users_assign_roles_excel.py excel-working/User_Assign_Role.xlsx
    -> Outputs: excel-working/User_Assign_Role_converted.csv
"""

import os
import sys
import csv
import openpyxl

# --- Safe Console Output on Windows ---
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')


def convert_users_assign_roles_excel(input_path):
    """
    Reads User Role Assignment matrix Excel file and converts to flat CSV format.
    Row 1: Role Names (Columns 2..N)
    Column 1: Username (Rows 2..N)
    Matrix Cells: 'Y', 'y', 'X', 'x' indicator flags
    """
    # --- Path & File Resolution ---
    if not os.path.splitext(input_path)[1]:
        input_path = input_path + ".xlsx"

    if not os.path.exists(input_path):
        print(f"[ERROR] Input file not found: {input_path}")
        sys.exit(1)

    dir_name = os.path.dirname(input_path)
    base_name = os.path.basename(input_path)
    stem, _ = os.path.splitext(base_name)

    out_csv_name = f"{stem}_converted.csv"
    out_csv_path = os.path.join(dir_name, out_csv_name) if dir_name else out_csv_name

    # --- Load Workbook ---
    try:
        wb_in = openpyxl.load_workbook(input_path, data_only=True)
    except Exception as e:
        print(f"[ERROR] Could not read Excel workbook '{input_path}': {e}")
        sys.exit(1)

    # Locate 'User Asignment' / 'Assignment' sheet or fallback to active sheet
    target_sheet = None
    for name in wb_in.sheetnames:
        clean_name = name.strip().lower()
        if 'user' in clean_name and ('asign' in clean_name or 'assign' in clean_name):
            target_sheet = wb_in[name]
            break

    if target_sheet is None:
        for name in wb_in.sheetnames:
            if 'assign' in name.strip().lower() or 'role' in name.strip().lower():
                target_sheet = wb_in[name]
                break

    if target_sheet is None:
        target_sheet = wb_in.active
        print(f"[INFO] Using active sheet '{target_sheet.title}'.")
    else:
        print(f"[INFO] Processing sheet: '{target_sheet.title}'")

    ws = target_sheet

    # --- Header Roles Extraction (Row 1, Cols 2..max_column) ---
    role_columns = []
    for col in range(2, ws.max_column + 1):
        raw_role = ws.cell(1, col).value
        if raw_role:
            role_name = str(raw_role).strip().upper()
            role_columns.append((col, role_name))

    if not role_columns:
        print(f"[ERROR] No role headers found in Row 1 (starting from Column 2).")
        sys.exit(1)

    print(f"[INFO] Detected {len(role_columns)} Role columns in header.")

    # --- Matrix Scanning & Flag Filtering ('Y', 'y', 'X', 'x') ---
    records = []
    users_with_roles = 0
    users_without_roles = 0
    empty_user_names = []
    unique_users = set()
    unique_roles_assigned = set()

    for row in range(2, ws.max_row + 1):
        raw_user = ws.cell(row, 1).value
        if not raw_user:
            continue

        username = str(raw_user).strip().upper()
        unique_users.add(username)
        found_flag = False

        for col, role_name in role_columns:
            val = str(ws.cell(row, col).value or '').strip().upper()

            # Filter criteria: accepts Y, y, X, x
            if val in ('Y', 'X'):
                records.append({
                    'username': username,
                    'role': role_name
                })
                unique_roles_assigned.add(role_name)
                found_flag = True

        # Fallback if no roles were marked for this user
        if not found_flag:
            records.append({
                'username': username,
                'role': 'N/A'
            })
            users_without_roles += 1
            empty_user_names.append(username)
        else:
            users_with_roles += 1

    # --- CSV Export (UTF-8 with BOM) ---
    try:
        with open(out_csv_path, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            writer.writerow(["Username", "Role"])
            for item in records:
                writer.writerow([item['username'], item['role']])
    except PermissionError:
        print(f"[ERROR] Permission denied when writing to '{out_csv_path}'. Please close the file in Excel and try again.")
        sys.exit(1)

    # --- Execution Summary ---
    total_users = users_with_roles + users_without_roles
    print(f"\n--- Execution Summary ---")
    print(f"[OK] Input File       : {input_path}")
    print(f"[OK] Total Users      : {total_users} ({users_with_roles} with roles, {users_without_roles} set to 'N/A')")
    if empty_user_names:
        print(f"[INFO] Empty Users    : {', '.join(empty_user_names)} -> assigned 'N/A'")
    print(f"[OK] Roles In Scope   : {len(role_columns)} available roles in matrix")
    print(f"[OK] Assigned Roles   : {len(unique_roles_assigned)} unique roles assigned")
    print(f"[OK] Output CSV       : {out_csv_path} ({len(records)} assignment rows)")


# --- CLI Entrypoint ---
if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] in ('-h', '--help'):
        print("Usage: python scripts/convert_users_assign_roles_excel.py [path_to_excel_file]")
        print("Example: python scripts/convert_users_assign_roles_excel.py excel-working/User_Assign_Role.xlsx")
        sys.exit(0)

    target_file = None

    if len(sys.argv) > 1:
        target_file = sys.argv[1]
    else:
        # Default candidates lookup
        candidates = [
            os.path.join("excel-working", "User_Assign_Role.xlsx"),
            os.path.join("excel-templates", "Template_User_Assign_Role.xlsx"),
            "User_Assign_Role.xlsx",
            "Users_Assign_Roles.xlsx",
            "User_Assign_Roles.xlsx"
        ]
        for cand in candidates:
            if os.path.exists(cand):
                target_file = cand
                break

    if not target_file:
        print("[ERROR] No input file provided or found in default locations.")
        print("Usage: python scripts/convert_users_assign_roles_excel.py <path_to_excel_file>")
        print("Example: python scripts/convert_users_assign_roles_excel.py excel-working/User_Assign_Role.xlsx")
        sys.exit(1)

    convert_users_assign_roles_excel(target_file)
