#!/usr/bin/env python3
"""
Convert Role Matrix Excel file to flat row format for SAP Role upload.
Filters Tcodes where cell value is 'Y'.
Outputs plain Excel sheet without formatting or index columns.

Usage:
    python scripts/convert_role_excel.py [input_excel_file]

Example:
    python scripts/convert_role_excel.py Role.xlsx
    -> Outputs: Role_converted.xlsx and Role_converted.csv
"""

import os
import sys
import csv
import openpyxl

# Ensure safe console output on Windows
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

def convert_role_excel(input_path):
    # Resolve input path and extension
    if not os.path.splitext(input_path)[1]:
        input_path = input_path + ".xlsx"

    if not os.path.exists(input_path):
        print(f"Error: File not found: {input_path}")
        sys.exit(1)

    # Derive output filenames: <original_name>_converted.xlsx / .csv
    dir_name = os.path.dirname(input_path)
    base_name = os.path.basename(input_path)
    stem, _ = os.path.splitext(base_name)

    out_excel_name = f"{stem}_converted.xlsx"
    out_csv_name = f"{stem}_converted.csv"

    out_excel_path = os.path.join(dir_name, out_excel_name) if dir_name else out_excel_name
    out_csv_path = os.path.join(dir_name, out_csv_name) if dir_name else out_csv_name

    wb_in = openpyxl.load_workbook(input_path, data_only=True)
    ws_in = wb_in.active

    roles_flat = []

    # Iterate over columns (each column starting from col 2 is a Role)
    for col in range(2, ws_in.max_column + 1):
        role_desc = str(ws_in.cell(1, col).value or '').strip()
        role_name = str(ws_in.cell(2, col).value or '').strip()
        if not role_name:
            continue

        for row in range(3, ws_in.max_row + 1):
            tcode = str(ws_in.cell(row, 1).value or '').strip()
            val = str(ws_in.cell(row, col).value or '').strip().upper()
            if val == 'Y' and tcode:
                roles_flat.append({
                    'role_name': role_name,
                    'role_desc': role_desc,
                    'tcode': tcode
                })

    # 1. Export Excel: Single sheet "Role_TCode_Rows", plain data, no formatting, no index column
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Role_TCode_Rows"

    # Header
    ws_out.append(["Role Name", "Role Description", "TCode"])

    # Data rows
    for item in roles_flat:
        ws_out.append([item['role_name'], item['role_desc'], item['tcode']])

    wb_out.save(out_excel_path)

    # 2. Export CSV for direct ABAP/text upload
    with open(out_csv_path, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        writer.writerow(["Role Name", "Role Description", "TCode"])
        for item in roles_flat:
            writer.writerow([item['role_name'], item['role_desc'], item['tcode']])

    print(f"[OK] Input        : {input_path}")
    print(f"[OK] Output Excel : {out_excel_path} (Sheet: 'Role_TCode_Rows', {len(roles_flat)} data rows)")
    print(f"[OK] Output CSV   : {out_csv_path} ({len(roles_flat)} data rows)")

if __name__ == "__main__":
    if len(sys.argv) > 1:
        target_file = sys.argv[1]
    else:
        # Default to Role.xlsx if present in current directory
        default_file = "Role.xlsx"
        if os.path.exists(default_file):
            target_file = default_file
        else:
            print("Usage: python scripts/convert_role_excel.py <path_to_excel_file>")
            print("Example: python scripts/convert_role_excel.py Role.xlsx")
            sys.exit(1)

    convert_role_excel(target_file)
