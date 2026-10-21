#!/usr/bin/env python3
"""
Convert Master Roles Matrix Excel file to flat row format for SAP Role upload.
Filters Transaction Codes (TCodes) where cell value is marked with an indicator ('Y', 'y', 'X', 'x').
If a role (column) has no indicator values, generates 1 row with TCode = 'N/A'.

Output columns:
1. Role Name
2. Role Description
3. TCode

Output:
- <filename>_converted.csv (UTF-8 with BOM for SAP GUI import)

Usage:
    python scripts/convert_role_01_master.py [input_excel_file]

Example:
    python scripts/convert_role_01_master.py excel-working/MM_Role_01_Master.xlsx
    -> Outputs: excel-working/converted/MM_Role_01_Master_converted.csv
"""

import os
import sys
import csv
import glob
import openpyxl

# --- Safe Console Output on Windows ---
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')


def is_indicator(val):
    """Checks if a cell value is an indicator flag or empty rather than a TCode."""
    if val is None:
        return True
    s = str(val).strip().upper()
    return s in ('Y', 'X', '0', '1', 'YES', 'NO', '', 'N/A', '-')


def convert_role_01_master(input_path, output_dir=None):
    """
    Reads Master Role matrix Excel file and converts to flat CSV format.
    Dynamically detects the TCode column and Master Role columns to support
    varied module layouts (e.g., FICO, PS, MM).
    Row 1: Role Description (preserves exact spacing)
    Row 2: Role Name
    Row 3+: TCodes (Detected column) and indicator flags ('Y', 'y', 'X', 'x')
    Output is saved to <parent_dir>/converted/ by default.
    """
    # --- Path & File Resolution ---
    if not os.path.splitext(input_path)[1]:
        input_path = input_path + ".xlsx"

    if not os.path.exists(input_path):
        print(f"[ERROR] Input file not found: {input_path}")
        return False

    base_name = os.path.basename(input_path)
    stem, _ = os.path.splitext(base_name)

    out_csv_name = f"{stem}_converted.csv"
    if output_dir is None:
        dir_name = os.path.dirname(input_path) or "."
        output_dir = os.path.join(dir_name, "converted")

    if output_dir.endswith('.csv'):
        out_csv_path = output_dir
        if os.path.dirname(out_csv_path):
            os.makedirs(os.path.dirname(out_csv_path), exist_ok=True)
    else:
        os.makedirs(output_dir, exist_ok=True)
        out_csv_path = os.path.join(output_dir, out_csv_name)

    # --- Load Workbook ---
    try:
        wb_in = openpyxl.load_workbook(input_path, data_only=True)
    except Exception as e:
        print(f"[ERROR] Could not read Excel workbook '{input_path}': {e}")
        return False

    # Locate target sheet: prefer sheet with 'MASTER' or 'LEADING' in title, else active sheet
    target_sheet = None
    for name in wb_in.sheetnames:
        low = name.strip().lower()
        if 'master' in low or 'leading' in low:
            target_sheet = wb_in[name]
            break

    if target_sheet is None:
        target_sheet = wb_in.active
        print(f"[INFO] Using active sheet '{target_sheet.title}'.")
    else:
        print(f"[INFO] Processing sheet: '{target_sheet.title}'")

    ws = target_sheet

    # --- 1. Dynamically Detect TCode Column ---
    best_tcode_col = 1
    max_tcode_count = 0

    for c in range(1, ws.max_column + 1):
        # Count non-indicator values in row 3+
        tcode_candidates = 0
        for r in range(3, ws.max_row + 1):
            v = ws.cell(r, c).value
            if v is not None and not is_indicator(v):
                tcode_candidates += 1

        if tcode_candidates > max_tcode_count:
            max_tcode_count = tcode_candidates
            best_tcode_col = c

    # --- 2. Dynamically Detect Master Role Columns ---
    role_cols = []
    for c in range(1, ws.max_column + 1):
        if c == best_tcode_col:
            continue
        rname = ws.cell(2, c).value
        if rname is not None and str(rname).strip():
            name_str = str(rname).strip()
            # Ignore label headers in template files
            if name_str.lower() in ('role name', 'role', 'tcode', 'transaction', 't-code'):
                continue
            rdesc = ws.cell(1, c).value
            role_desc = str(rdesc) if rdesc is not None else ''
            role_cols.append((c, name_str, role_desc))

    print(f"[INFO] TCode column detected: Col {best_tcode_col} ({max_tcode_count} TCodes). Master roles: {len(role_cols)} columns.")

    roles_flat = []
    roles_with_tcodes = 0
    roles_without_tcodes = 0
    no_flag_role_names = []

    # --- 3. Matrix Scanning & Flag Filtering ('Y', 'y', 'X', 'x') ---
    for col, role_name, role_desc in role_cols:
        found_flag = False
        for row in range(3, ws.max_row + 1):
            tcode = str(ws.cell(row, best_tcode_col).value or '').strip()
            val = str(ws.cell(row, col).value or '').strip().upper()

            # Filter criteria: accepts Y, y, X, x
            if val in ('Y', 'X') and tcode:
                roles_flat.append({
                    'role_name': role_name,
                    'role_desc': role_desc,
                    'tcode': tcode
                })
                found_flag = True

        # If no indicator ('Y'/'X') was found for this role column, create 1 fallback row with TCode = 'N/A'
        if not found_flag:
            roles_flat.append({
                'role_name': role_name,
                'role_desc': role_desc,
                'tcode': 'N/A'
            })
            roles_without_tcodes += 1
            no_flag_role_names.append(role_name)
        else:
            roles_with_tcodes += 1

    # --- CSV Export (UTF-8 with BOM) ---
    try:
        with open(out_csv_path, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            writer.writerow(["Role Name", "Role Description", "TCode"])
            for item in roles_flat:
                writer.writerow([item['role_name'], item['role_desc'], item['tcode']])
    except PermissionError:
        print(f"[ERROR] Permission denied when writing to '{out_csv_path}'. Please close the file in Excel and try again.")
        return False

    # --- Execution Summary ---
    total_roles = roles_with_tcodes + roles_without_tcodes
    print(f"\n--- Execution Summary ---")
    print(f"[OK] Input File       : {input_path}")
    print(f"[OK] Total Roles      : {total_roles} ({roles_with_tcodes} with TCodes, {roles_without_tcodes} set to 'N/A')")
    if no_flag_role_names:
        print(f"[INFO] Empty Roles    : {', '.join(no_flag_role_names)} -> assigned 'N/A'")
    print(f"[OK] Output CSV       : {out_csv_path} ({len(roles_flat)} data rows)")
    return True


# Compatibility alias
convert_role_master_excel = convert_role_01_master


# --- CLI Entrypoint ---
if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] in ('-h', '--help'):
        print("Usage: python scripts/convert_role_01_master.py [path_to_excel_file]")
        print("Example: python scripts/convert_role_01_master.py excel-working/MM_Role_01_Master.xlsx")
        sys.exit(0)

    target_files = []

    if len(sys.argv) > 1:
        for arg in sys.argv[1:]:
            matches = [m for m in glob.glob(arg) if not os.path.basename(m).startswith('~$')]
            if matches:
                target_files.extend(matches)
            elif not os.path.basename(arg).startswith('~$'):
                target_files.append(arg)
    else:
        # Default candidates lookup in working and template directories
        raw_candidates = (glob.glob(os.path.join("excel-working", "*_01_Master.xlsx")) +
                          glob.glob(os.path.join("excel-working", "*_Master.xlsx")))
        unique_working = sorted(list({c for c in raw_candidates if not os.path.basename(c).startswith('~$')}))
        
        if unique_working:
            target_files = unique_working
        else:
            raw_templates = (glob.glob(os.path.join("excel-templates", "*_01_Master.xlsx")) +
                             glob.glob(os.path.join("excel-templates", "*_Master.xlsx")))
            unique_templates = sorted(list({c for c in raw_templates if not os.path.basename(c).startswith('~$')}))
            if unique_templates:
                target_files = unique_templates
            else:
                target_files = ["Template_Role_01_Master.xlsx"]

    if not target_files:
        print("[ERROR] No input file provided or found in default locations.")
        print("Usage: python scripts/convert_role_01_master.py <path_to_excel_file>")
        print("Example: python scripts/convert_role_01_master.py excel-working/MM_Role_01_Master.xlsx")
        sys.exit(1)

    success = True
    for tf in target_files:
        print(f"\n==================================================")
        print(f"[RUNNING] Converting Master Role: {tf}")
        print(f"==================================================")
        if not convert_role_01_master(tf):
            success = False

    if not success:
        sys.exit(1)
