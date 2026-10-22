#!/usr/bin/env python3
"""
Convert Organizational Levels and Authorization Objects & Activities from Excel matrix
into two separate flat CSV files for SAP Role upload.

Features:
- Reads consolidated role workbooks (Sheet 2: Derived role, Org, and Object data)
  or standalone matrix files.
- Automatically classifies columns:
  * Organizational Levels (matched against Org_Levels.csv / SAP USORG) -> outputs <stem>_03a_Org_Levels_converted.csv
  * Authorization Objects / Fields (matched against Object_Levels.csv)  -> outputs <stem>_03b_Auth_Objects_converted.csv
- Strictly non-overlapping: Org values are maintained in AGR_1252 via Org Level variables ($VARBL)
  and are NEVER mixed into Object output (AGR_1251).
- Excludes TCode check (S_TCODE-TCD) as it is maintained via the role menu.
- Supports multi-object columns where one value applies to multiple objects (e.g. Person Responsible in C_PROJ_VNR and C_PRPS_VNR).

Output formats:
1. Org CSV (03a_Org_Levels):
   Col 1: Role
   Col 2: Org_Level ($VARBL)
   Col 3: Value
2. Object CSV (03b_Auth_Objects):
   Col 1: Role
   Col 2: object-field (e.g., F_BKPF_BUK-BUKRS)
   Col 3: value

Outputs:
- <parent_dir>/converted/<stem>_03a_Org_Levels_converted.csv
- <parent_dir>/converted/<stem>_03b_Auth_Objects_converted.csv
"""

import os
import sys
import csv
import glob
import re
import openpyxl

# --- Safe Console Output on Windows ---
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

# Standard SAP Org Variable Synonyms & Common Business Terms
COMMON_ORG_SYNONYMS = {
    'company code': '$BUKRS',
    'controlling area': '$KOKRS',
    'co area': '$KOKRS',
    'operating concern': '$ERKRS',
    'plant': '$WERKS',
    'plant / valuation area': '$WERKS',
    'valuation area': '$BWKEY',
    'profit center': '$PRCTR',
    'profit centers': '$PRCTR',
    'purchasing org': '$EKORG',
    'purchasing organization': '$EKORG',
    'purchasing group': '$EKGRP',
    'sales org': '$VKORG',
    'sales organization': '$VKORG',
    'distribution channel': '$VTWEG',
    'shipping point': '$VSTEL',
    'division': '$SPART',
    'business area': '$GSBER',
    'functional area': '$FKBER',
    'credit control area': '$KKBER',
    'storage type': '$LGTYP',
    'warehouse number': '$LGNUM',
    'warehouse number / warehouse complex': '$LGNUM',
    'sales office': '$VKBUR',
    'sales group': '$VKGRP',
    'work center': '$ARBPL',
}

# Common Business Field to SAP Authorization Object-Field Mapping
COMMON_OBJECT_FIELDS = {
    'segment': ['F_FAGL_SEG-SEGMENT'],
    'business place': ['F_BUK_BUPL-BUPLA'],
    'cost center': ['K_CSKS-KOSTL'],
    'cost centers': ['K_CSKS-KOSTL'],
    'order type': ['K_ORDER-AUFART'],
    'order types': ['K_ORDER-AUFART'],
    'asset class': ['A_S_ANLKL-ANLKL'],
    'credit segment': ['F_UKM_SGMT-UKM_SGMNT'],
    'costing variant': ['K_KEKO-KLVAR'],
    'person responsible': ['C_PROJ_VNR-PS_VERNR', 'C_PRPS_VNR-PS_VERNR'],
    'project type': ['C_PRPS_ART-PS_PRART'],
    'release group': ['M_EINK_FRG-FRGGR'],
    'release code': ['M_EINK_FRG-FRGCO'],
    'storage location': ['M_MSEG_LGO-LGORT'],
}


def load_org_levels_reference(org_levels_path):
    """Loads Org_Levels.csv mapping dictionary."""
    if not os.path.exists(org_levels_path):
        return {}

    mapping = {}
    try:
        with open(org_levels_path, 'r', encoding='utf-8-sig', errors='replace') as f:
            reader = csv.DictReader(f)
            for row in reader:
                varbl = row.get('Org Level Variable', '').strip().upper()
                desc = row.get('Org Level Description', '').strip().lower()
                fld = row.get('Field Name', '').strip().upper()

                if desc and varbl:
                    mapping[desc] = varbl
                if fld and varbl:
                    mapping[fld] = varbl
                    mapping[fld.lstrip('/')] = varbl
                if varbl:
                    mapping[varbl] = varbl
                    mapping[varbl.lstrip('$')] = varbl
    except Exception as e:
        print(f"[WARN] Could not load Org_Levels.csv: {e}")

    return mapping


def load_object_levels_reference(obj_levels_path):
    """Loads Object_Levels.csv into lookup structures."""
    if not os.path.exists(obj_levels_path):
        return {}, {}, {}

    pair_dict = {}
    object_fields_map = {}
    desc_dict = {}

    try:
        with open(obj_levels_path, 'r', encoding='utf-8-sig', errors='replace') as f:
            reader = csv.DictReader(f)
            for row in reader:
                pair = row.get('Object-Field', '').strip().upper()
                obj = row.get('Object', '').strip().upper()
                field = row.get('Field', '').strip().upper()
                fdesc = row.get('Field Description', '').strip()
                odesc = row.get('Object Description', '').strip()

                if pair:
                    pair_dict[pair] = (obj, field, fdesc, odesc)

                if obj:
                    if obj not in object_fields_map:
                        object_fields_map[obj] = []
                    object_fields_map[obj].append((field, fdesc, odesc, pair))

                if fdesc:
                    desc_dict[fdesc.lower()] = pair
    except Exception as e:
        print(f"[WARN] Could not load Object_Levels.csv: {e}")

    return pair_dict, object_fields_map, desc_dict


def classify_column(header_text, org_ref_map, obj_pairs, obj_fields_map, obj_desc_dict):
    """
    Classifies an Excel column header as either:
      ('ORG', [varbl])
      ('OBJECT', [object_field_1, object_field_2, ...])
      ('SKIP_TCODE', 'S_TCODE-TCD')
      or None if unclassified.

    Guarantees strict non-overlapping classification:
    Org levels are NEVER routed to Object CSV.
    """
    raw = str(header_text or '').strip()
    if not raw:
        return None

    clean = raw.replace('\n', ' ').strip()
    clean_lower = clean.lower()

    # =========================================================================
    # STEP 0: EXCLUDE TCODES (S_TCODE-TCD)
    # =========================================================================
    if any(k in clean_lower for k in ('s_tcode', 's_tcode-tcd', 'transaction code check')):
        return ('SKIP_TCODE', 'S_TCODE-TCD')

    # =========================================================================
    # STEP 1: CHECK IF ORGANIZATIONAL LEVEL
    # =========================================================================
    # 1.1 Direct parenthesized variable in header, e.g., ($BUKRS), ($CEERKRS), ($KOKRS)
    m_var = re.search(r'\(\$([A-Za-z0-9_]+)\)', raw)
    if m_var:
        return ('ORG', [f"${m_var.group(1).upper()}"])

    # 1.2 Standalone $VARBL in header
    m_dollar = re.search(r'\$([A-Za-z0-9_]+)', raw)
    if m_dollar:
        cand_var = f"${m_dollar.group(1).upper()}"
        if cand_var in org_ref_map.values() or cand_var in COMMON_ORG_SYNONYMS.values():
            return ('ORG', [cand_var])

    # 1.3 Check parenthesized field name if it is an Org Level (e.g. (BUKRS), (WERKS), (KOKRS))
    m_paren = re.search(r'\((.*?)\)', clean)
    if m_paren:
        inside = m_paren.group(1).strip()
        tokens = [t.strip().upper() for t in inside.split(',') if t.strip()]
        for t in tokens:
            t_clean = t.lstrip('$')
            if t in org_ref_map:
                return ('ORG', [org_ref_map[t]])
            if t_clean in org_ref_map:
                return ('ORG', [org_ref_map[t_clean]])

    # 1.4 Check common business synonyms
    stripped_paren = re.sub(r'\(.*?\)', '', clean).strip().lower()
    if stripped_paren in COMMON_ORG_SYNONYMS:
        return ('ORG', [COMMON_ORG_SYNONYMS[stripped_paren]])
    if clean_lower in COMMON_ORG_SYNONYMS:
        return ('ORG', [COMMON_ORG_SYNONYMS[clean_lower]])

    # 1.5 Check Org_Levels.csv reference
    if stripped_paren in org_ref_map:
        return ('ORG', [org_ref_map[stripped_paren]])
    if clean_lower in org_ref_map:
        return ('ORG', [org_ref_map[clean_lower]])

    # 1.6 Check if header is an explicit OBJECT-FIELD whose field is an Org Level (e.g. F_BKPF_BUK-BUKRS)
    m_pair_std = re.search(r'([A-Za-z0-9_]{3,10})[-:.]([A-Za-z0-9_]{2,10})', raw)
    if m_pair_std:
        fld_part = m_pair_std.group(2).upper()
        if fld_part in org_ref_map:
            return ('ORG', [org_ref_map[fld_part]])

    # =========================================================================
    # STEP 2: CHECK IF AUTHORIZATION OBJECT / FIELD (Non-Org)
    # =========================================================================
    # 2.1 Check common business object fields mapping first (e.g. Storage Location -> M_MSEG_LGO-LGORT)
    if stripped_paren in COMMON_OBJECT_FIELDS:
        return ('OBJECT', COMMON_OBJECT_FIELDS[stripped_paren])
    if clean_lower in COMMON_OBJECT_FIELDS:
        return ('OBJECT', COMMON_OBJECT_FIELDS[clean_lower])

    # 2.2 Check parenthesized content: could be single or comma-separated objects/fields
    if m_paren:
        inside = m_paren.group(1).strip()
        tokens = [t.strip().upper() for t in inside.split(',') if t.strip()]
        matched_pairs = []

        for t in tokens:
            # Token is an explicit OBJECT-FIELD (e.g. F_BKPF_BUK-BUKRS or F_BKPF_BUK:BUKRS)
            if '-' in t or ':' in t or '.' in t:
                norm_pair = re.sub(r'[:.]', '-', t)
                fld_cand = norm_pair.split('-')[1] if '-' in norm_pair else ''
                # Protect: If field belongs to Org Levels, redirect to Org Level
                if fld_cand in org_ref_map:
                    return ('ORG', [org_ref_map[fld_cand]])
                matched_pairs.append(norm_pair)
                continue

            # Token is an OBJECT name (e.g. K_KEKO, C_PROJ_VNR, C_PRPS_VNR)
            if t in obj_fields_map:
                fields = obj_fields_map[t]
                chosen = None
                # First check if field description or field name matches header
                for fld, fdesc, odesc, p in fields:
                    if fdesc.lower() in clean_lower or fld.lower() in clean_lower:
                        chosen = p
                        break
                if not chosen:
                    # If header mentions activity/actvt, pick ACTVT field
                    if 'activity' in clean_lower or 'actvt' in clean_lower:
                        actvt_flds = [p for f, d, o, p in fields if 'actvt' in f.lower()]
                        chosen = actvt_flds[0] if actvt_flds else fields[0][3]
                    else:
                        # Pick non-ACTVT field
                        non_actvt = [p for f, d, o, p in fields if 'actvt' not in f.lower()]
                        chosen = non_actvt[0] if non_actvt else fields[0][3]
                if chosen and chosen not in matched_pairs:
                    matched_pairs.append(chosen)

            # Token is a FIELD name (e.g. FRGGR, FRGCO, LGORT)
            # Check if any obj_pairs end with -TOKEN
            matching_pairs = [p for p in obj_pairs if p.endswith(f"-{t}")]
            if matching_pairs:
                chosen_pair = None
                for p in matching_pairs:
                    obj_name, fld_name, fdesc, odesc = obj_pairs[p]
                    if any(w in clean_lower for w in fdesc.lower().split()) or any(w in clean_lower for w in odesc.lower().split()):
                        chosen_pair = p
                        break
                if not chosen_pair:
                    chosen_pair = matching_pairs[0]
                if chosen_pair not in matched_pairs:
                    matched_pairs.append(chosen_pair)

        if matched_pairs:
            return ('OBJECT', matched_pairs)

    # 2.2 Check standalone OBJECT-FIELD in header
    if m_pair_std:
        fld_part = m_pair_std.group(2).upper()
        if fld_part in org_ref_map:
            return ('ORG', [org_ref_map[fld_part]])
        pair = f"{m_pair_std.group(1).upper()}-{m_pair_std.group(2).upper()}"
        if pair in obj_pairs:
            return ('OBJECT', [pair])

    # 2.3 Check common business object fields mapping
    if stripped_paren in COMMON_OBJECT_FIELDS:
        return ('OBJECT', COMMON_OBJECT_FIELDS[stripped_paren])
    if clean_lower in COMMON_OBJECT_FIELDS:
        return ('OBJECT', COMMON_OBJECT_FIELDS[clean_lower])

    # 2.4 Check Object_Levels.csv field description dictionary
    if stripped_paren in obj_desc_dict:
        return ('OBJECT', [obj_desc_dict[stripped_paren]])
    if clean_lower in obj_desc_dict:
        return ('OBJECT', [obj_desc_dict[clean_lower]])

    return None


def convert_role_03_org_obj(input_path, output_dir=None):
    """
    Reads an Excel file and converts Organizational levels and Authorization Objects
    into two separate flat CSV files:
      - <stem>_03a_Org_Levels_converted.csv (Role, Org_Level, Value) -> for AGR_1252
      - <stem>_03b_Auth_Objects_converted.csv (Role, object-field, value) -> for AGR_1251
    """
    if not os.path.splitext(input_path)[1]:
        input_path = input_path + ".xlsx"

    if not os.path.exists(input_path):
        print(f"[ERROR] Input file not found: {input_path}")
        return False

    base_name = os.path.basename(input_path)
    stem, _ = os.path.splitext(base_name)

    # Resolve output directory
    if output_dir is None:
        dir_name = os.path.dirname(input_path) or "."
        output_dir = os.path.join(dir_name, "converted")
    os.makedirs(output_dir, exist_ok=True)

    # Determine output file stems
    clean_stem = re.sub(r'_(03[ab]?_)?(org|object|obj).*$', '', stem, flags=re.IGNORECASE)
    if not clean_stem.endswith('_Role') and stem.endswith('_Role'):
        clean_stem = stem
    elif not clean_stem:
        clean_stem = stem

    out_org_path = os.path.join(output_dir, f"{clean_stem}_03a_Org_Levels_converted.csv")
    out_obj_path = os.path.join(output_dir, f"{clean_stem}_03b_Auth_Objects_converted.csv")

    # Load Reference Dictionaries
    org_ref_path = os.path.join("excel-templates", "Org_Levels.csv")
    obj_ref_path = os.path.join("excel-templates", "Object_Levels.csv")

    org_ref_map = load_org_levels_reference(org_ref_path)
    obj_pairs, obj_fields_map, obj_desc_dict = load_object_levels_reference(obj_ref_path)

    # Load Workbook
    try:
        wb_in = openpyxl.load_workbook(input_path, data_only=True)
    except Exception as e:
        print(f"[ERROR] Could not read Excel workbook '{input_path}': {e}")
        return False

    # Sheet Selection:
    # Rule: For consolidated role files (e.g. FICO_Role.xlsx), Sheet 2 (Index 1) is Derived + Org + Object.
    # Otherwise, prefer sheet containing 'derive', 'org', 'object', or active/Sheet 1.
    is_consolidated = not any(k in stem.lower() for k in ('_01_', '_02_', '_03_', '_04_', '_master', '_derive', '_org', '_object', '_composite'))

    target_sheet = None
    if is_consolidated and len(wb_in.sheetnames) >= 2:
        target_sheet = wb_in.worksheets[1]
        print(f"[INFO] Consolidated file: Using Sheet 2 (Index 1): '{target_sheet.title}' for Derived Role & Org/Object data.")
    else:
        for name in wb_in.sheetnames:
            low = name.strip().lower()
            if any(k in low for k in ('org', 'object', 'derive')):
                target_sheet = wb_in[name]
                break
        if target_sheet is None:
            target_sheet = wb_in.worksheets[0] if wb_in.worksheets else wb_in.active
        print(f"[INFO] Processing sheet: '{target_sheet.title}'")

    ws = target_sheet

    # Detect Role Name Column:
    # In Sheet 2 of consolidated file: Col 1 = Master, Col 2 = Derived Role, Col 3 = Desc, Col 4+ = Data
    role_col = 2
    data_start_col = 4

    sample_col1 = [str(ws.cell(r, 1).value or '').strip() for r in range(2, min(ws.max_row + 1, 10))]
    sample_col2 = [str(ws.cell(r, 2).value or '').strip() for r in range(2, min(ws.max_row + 1, 10))]

    if any(s.startswith(('DR', 'Z_')) for s in sample_col2):
        role_col = 2
        data_start_col = 4
    elif any(s.startswith(('DR', 'Z_')) for s in sample_col1):
        role_col = 1
        data_start_col = 2
    else:
        h1 = str(ws.cell(1, 1).value or '').lower()
        h2 = str(ws.cell(1, 2).value or '').lower()
        if 'derive' in h2 or 'role' in h2:
            role_col = 2
            data_start_col = 4 if ws.max_column >= 4 else 3
        else:
            role_col = 1
            data_start_col = 2

    print(f"[INFO] Role column: Col {role_col}, Org/Object data columns start at: Col {data_start_col}")

    # Classify Columns
    classified_cols = []
    org_col_count = 0
    obj_col_count = 0

    header_row = 1
    for c in range(data_start_col, ws.max_column + 1):
        h_val = ws.cell(header_row, c).value
        if h_val is None:
            continue
        h_text = str(h_val).strip()
        if not h_text:
            continue

        kind_info = classify_column(h_text, org_ref_map, obj_pairs, obj_fields_map, obj_desc_dict)
        if kind_info:
            kind, idents = kind_info
            if kind == 'SKIP_TCODE':
                print(f"[INFO] Col {c:2d}: '{h_text}' (S_TCODE-TCD) -> Skipped from authorization convert (managed via Menu).")
                continue

            classified_cols.append((c, kind, idents, h_text))
            if kind == 'ORG':
                org_col_count += len(idents)
            elif kind == 'OBJECT':
                obj_col_count += len(idents)
        else:
            print(f"[WARN] Col {c:2d} '{h_text}' could not be classified as Org Level or Authorization Object. Skipped.")

    print(f"[INFO] Classification summary: {org_col_count} Org Level mappings, {obj_col_count} Auth Object/Field mappings.")
    for c, kind, idents, h_text in classified_cols:
        clean_name = h_text.replace('\n', ' ')
        ident_str = ', '.join(idents)
        print(f"  - Col {c:2d}: [{kind:6s}] {clean_name:35s} -> {ident_str}")

    # Process Data Rows
    org_rows = []
    obj_rows = []
    roles_processed = 0

    for r in range(header_row + 1, ws.max_row + 1):
        role_val = ws.cell(r, role_col).value
        if role_val is None:
            continue
        role_name = str(role_val).strip()
        if not role_name or role_name.lower() in ('none', 'role', 'derived role', 'role name'):
            continue

        roles_processed += 1

        for c, kind, idents, _ in classified_cols:
            val_cell = ws.cell(r, c).value
            if val_cell is None:
                continue

            if isinstance(val_cell, float) and val_cell.is_integer():
                val_str = str(int(val_cell))
            else:
                val_str = str(val_cell).strip()

            if not val_str or val_str.lower() in ('none', 'n/a', '-'):
                continue

            # Populate rows
            if kind == 'ORG':
                for varbl in idents:
                    org_rows.append([role_name, varbl, val_str])
            elif kind == 'OBJECT':
                for obj_field in idents:
                    obj_rows.append([role_name, obj_field, val_str])

    # Write Org CSV
    try:
        with open(out_org_path, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            writer.writerow(["Role", "Org_Level", "Value"])
            writer.writerows(org_rows)
        print(f"[OK] Wrote Org Levels CSV   : {out_org_path} ({len(org_rows)} rows across {roles_processed} roles)")
    except Exception as e:
        print(f"[ERROR] Could not write Org Levels CSV '{out_org_path}': {e}")
        return False

    # Write Object CSV
    try:
        with open(out_obj_path, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            writer.writerow(["Role", "object-field", "value"])
            writer.writerows(obj_rows)
        print(f"[OK] Wrote Auth Objects CSV: {out_obj_path} ({len(obj_rows)} rows across {roles_processed} roles)")
    except Exception as e:
        print(f"[ERROR] Could not write Auth Objects CSV '{out_obj_path}': {e}")
        return False

    return True


# --- CLI Entrypoint ---
if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] in ('-h', '--help'):
        print("Usage: python scripts/convert_role_03_org_obj.py [path_to_excel_file]")
        print("Example: python scripts/convert_role_03_org_obj.py excel-working/FICO_Role.xlsx")
        sys.exit(0)

    target_files = []

    if len(sys.argv) > 1:
        for arg in sys.argv[1:]:
            matches = [m for m in glob.glob(arg) if not os.path.basename(m).startswith('~$')]
            if matches:
                target_files.extend(matches)
            elif not os.path.basename(arg).startswith('~$'):
                target_files.append(arg)
    else:
        # Default lookup in excel-working: prioritize consolidated *Role.xlsx, include remaining legacy
        working_consolidated = [f for f in glob.glob(os.path.join("excel-working", "*Role*.xlsx"))
                                if not os.path.basename(f).startswith('~$')
                                and not any(k in os.path.basename(f) for k in ('_01_', '_02_', '_03_', '_04_', '_converted'))]
        working_legacy = [f for f in glob.glob(os.path.join("excel-working", "*_03_*.xlsx"))
                          if not os.path.basename(f).startswith('~$')]

        consolidated_stems = {re.sub(r'_?Role.*$', '', os.path.basename(f), flags=re.I).upper() for f in working_consolidated}
        filtered_legacy = [f for f in working_legacy if re.sub(r'_?Role.*$', '', os.path.basename(f), flags=re.I).upper() not in consolidated_stems]

        unique_working = sorted(working_consolidated + filtered_legacy)

        if unique_working:
            target_files = unique_working
        else:
            tpl_consolidated = [f for f in glob.glob(os.path.join("excel-templates", "*Role*.xlsx"))
                                if not os.path.basename(f).startswith('~$')
                                and not any(k in os.path.basename(f) for k in ('_01_', '_02_', '_03_', '_04_', '_converted'))]
            tpl_legacy = [f for f in glob.glob(os.path.join("excel-templates", "*_03_*.xlsx"))
                          if not os.path.basename(f).startswith('~$')]
            tpl_stems = {re.sub(r'_?Role.*$', '', os.path.basename(f), flags=re.I).upper() for f in tpl_consolidated}
            filtered_tpl = [f for f in tpl_legacy if re.sub(r'_?Role.*$', '', os.path.basename(f), flags=re.I).upper() not in tpl_stems]

            unique_templates = sorted(tpl_consolidated + filtered_tpl)
            if unique_templates:
                target_files = unique_templates
            else:
                target_files = ["Template_Role.xlsx"]

    if not target_files:
        print("[ERROR] No input file provided or found in default locations.")
        print("Usage: python scripts/convert_role_03_org_obj.py <path_to_excel_file>")
        sys.exit(1)

    success = True
    for tf in target_files:
        print(f"\n==================================================")
        print(f"[RUNNING] Converting Org & Object Levels: {tf}")
        print(f"==================================================")
        if not convert_role_03_org_obj(tf):
            success = False

    if not success:
        sys.exit(1)
