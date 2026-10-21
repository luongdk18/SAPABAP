#!/usr/bin/env python3
"""
Convert Master Role Matrix Excel file to flat row format for SAP Role upload.
Filters Tcodes where cell value is 'Y'.
If a role (column) has no 'Y' values, creates 1 row with TCode = 'N/A'.
Outputs CSV file (_converted.csv) for SAP import.

Usage:
    python scripts/convert_role_master_excel.py [input_excel_file]

Example:
    python scripts/convert_role_master_excel.py Roles_Master.xlsx
    -> Outputs: Roles_Master_converted.csv
"""

import os
import sys
import csv
import openpyxl

# Ensure safe console output on Windows
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

def convert_role_master_excel(input_path):
    # Resolve input path and extension
    if not os.path.splitext(input_path)[1]:
        input_path = input_path + ".xlsx"

    if not os.path.exists(input_path):
        print(f"Error: File not found: {input_path}")
        sys.exit(1)

    # Derive output filename: <original_name>_converted.csv
    dir_name = os.path.dirname(input_path)
    base_name = os.path.basename(input_path)
    stem, _ = os.path.splitext(base_name)

    out_csv_name = f"{stem}_converted.csv"
    out_csv_path = os.path.join(dir_name, out_csv_name) if dir_name else out_csv_name

    wb_in = openpyxl.load_workbook(input_path, data_only=True)
    ws_in = wb_in.active

    roles_flat = []
    roles_with_tcodes = 0
    roles_without_tcodes = 0
    no_y_role_names = []

    # Iterate over columns (each column starting from col 2 is a Role)
    for col in range(2, ws_in.max_column + 1):
        cell_desc_val = ws_in.cell(1, col).value
        role_desc = str(cell_desc_val) if cell_desc_val is not None else ''
        role_name = str(ws_in.cell(2, col).value or '').strip()
        if not role_name:
            continue

        found_y = False
        for row in range(3, ws_in.max_row + 1):
            tcode = str(ws_in.cell(row, 1).value or '').strip()
            val = str(ws_in.cell(row, col).value or '').strip().upper()
            if val == 'Y' and tcode:
                roles_flat.append({
                    'role_name': role_name,
                    'role_desc': role_desc,
                    'tcode': tcode
                })
                found_y = True

        # If no 'Y' was found for this role column, create 1 row with TCode = 'N/A'
        if not found_y:
            roles_flat.append({
                'role_name': role_name,
                'role_desc': role_desc,
                'tcode': 'N/A'
            })
            roles_without_tcodes += 1
            no_y_role_names.append(role_name)
        else:
            roles_with_tcodes += 1

    # Export CSV for direct ABAP/text upload (UTF-8 with BOM)
    try:
        with open(out_csv_path, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            writer.writerow(["Role Name", "Role Description", "TCode"])
            for item in roles_flat:
                writer.writerow([item['role_name'], item['role_desc'], item['tcode']])
    except PermissionError:
        print(f"[ERROR] Permission denied when writing to '{out_csv_path}'. Please close the file and try again.")
        sys.exit(1)

    total_roles = roles_with_tcodes + roles_without_tcodes
    print(f"[OK] Input        : {input_path}")
    print(f"[OK] Total Roles  : {total_roles} ({roles_with_tcodes} with TCodes, {roles_without_tcodes} without TCodes set to 'N/A')")
    if no_y_role_names:
        print(f"[INFO] Roles with TCode 'N/A': {', '.join(no_y_role_names)}")
    print(f"[OK] Output CSV   : {out_csv_path} ({len(roles_flat)} data rows)")

if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] in ('-h', '--help'):
        print("Usage: python scripts/convert_role_master_excel.py [path_to_excel_file]")
        print("Example: python scripts/convert_role_master_excel.py Roles_Master.xlsx")
        sys.exit(0)

    if len(sys.argv) > 1:
        target_file = sys.argv[1]
    else:
        # Default lookup in current directory / working folders / templates
        candidates = [
            os.path.join("excel-working", "FICO_Roles_Master.xlsx"),
            os.path.join("excel-templates", "Template_Roles_Master.xlsx"),
            "Roles_Master.xlsx",
            "FICO_Roles_Master.xlsx",
            "Role.xlsx"
        ]
        target_file = None
        for cand in candidates:
            if os.path.exists(cand):
                target_file = cand
                break

        if not target_file:
            print("Usage: python scripts/convert_role_master_excel.py <path_to_excel_file>")
            print("Example: python scripts/convert_role_master_excel.py Roles_Master.xlsx")
            sys.exit(1)

    convert_role_master_excel(target_file)
