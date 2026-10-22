#!/usr/bin/env python3
"""
Generate standardized, masked template files for SAP Role and User upload in 'excel-templates/'.

Generated Files:
1. Template_Role.xlsx (Consolidated 3-sheet Role Workbook)
   - Sheet 1: Master Roles (Menu & TCodes)
   - Sheet 2: Derived Roles & Org/Object Data
   - Sheet 3: Composite Roles
   -> Auto-converts to:
      * Template_Role_01_Master_converted.csv
      * Template_Role_02_Derived_converted.csv
      * Template_Role_03a_Org_Levels_converted.csv
      * Template_Role_03b_Auth_Objects_converted.csv
      * Template_Role_04_Composite_converted.csv
2. Template_User_01_Create.xlsx -> Template_User_01_Create_converted.csv
3. Template_User_02_Assign.xlsx -> Template_User_02_Assign_converted.csv

Usage:
    python scripts/create_templates.py
"""

import os
import sys
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# --- Safe Console Output on Windows ---
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

# --- Directory & Style Definitions ---
OUT_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "excel-templates")
CONVERTED_DIR = os.path.join(OUT_DIR, "converted")
os.makedirs(OUT_DIR, exist_ok=True)
os.makedirs(CONVERTED_DIR, exist_ok=True)

header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
role_name_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
role_name_font = Font(name="Calibri", size=11, bold=True, color="000000")
data_font = Font(name="Calibri", size=11)
center_align = Alignment(horizontal="center", vertical="center")
left_align = Alignment(horizontal="left", vertical="center")
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)


# --- 1. Consolidated Role Template (Template_Role.xlsx) ---
def create_consolidated_role_template():
    wb = openpyxl.Workbook()

    # -------------------------------------------------------------
    # Sheet 1: Master Roles
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Master"

    roles = [
        {"name": "Z_FI_PERIOD_CONTROL", "desc": "GL Period Control", "tcodes": ["OB52", "S_ALR_87003642"]},
        {"name": "Z_FI_GL_MASTER_MAINT", "desc": "GL Master Data Maintenance", "tcodes": ["FS00", "FSP0", "FSS0"]},
        {"name": "Z_FI_JOURNAL_POST", "desc": "General Ledger Journal Posting", "tcodes": ["FB01", "FB02", "FB03", "F-02", "FB50"]},
        {"name": "Z_FI_JOURNAL_DISPLAY", "desc": "General Ledger Journal Display", "tcodes": ["FB03", "FBL3N", "FS10N"]},
        {"name": "Z_FI_AP_INVOICE_PROCESS", "desc": "Accounts Payable Invoice Processing", "tcodes": ["FB60", "FB02", "FB03", "F-53"]},
        {"name": "Z_FI_GL_REPORTS_EMPTY", "desc": "General Ledger Custom Reporting", "tcodes": []}
    ]

    all_tcodes = sorted(list({tc for r in roles for tc in r["tcodes"]}))
    c_hdr = ws1.cell(1, 1, value="TCode / Role")
    c_hdr.font = header_font
    c_hdr.fill = header_fill
    c_hdr.border = thin_border

    for col_idx, r in enumerate(roles, start=2):
        c_desc = ws1.cell(1, col_idx, value=r["desc"])
        c_desc.font = header_font
        c_desc.fill = header_fill
        c_desc.border = thin_border

        c_name = ws1.cell(2, col_idx, value=r["name"])
        c_name.font = role_name_font
        c_name.fill = role_name_fill
        c_name.border = thin_border

    for row_idx, tc in enumerate(all_tcodes, start=3):
        c_tc = ws1.cell(row_idx, 1, value=tc)
        c_tc.font = role_name_font
        c_tc.fill = role_name_fill
        c_tc.border = thin_border

        for col_idx, r in enumerate(roles, start=2):
            val = "X" if tc in r["tcodes"] else ""
            c_val = ws1.cell(row_idx, col_idx, value=val)
            c_val.font = data_font
            c_val.alignment = center_align
            c_val.border = thin_border

    ws1.column_dimensions['A'].width = 18
    for col_letter in ['B', 'C', 'D', 'E', 'F', 'G']:
        ws1.column_dimensions[col_letter].width = 30

    # -------------------------------------------------------------
    # Sheet 2: Derived & Org + Object Levels
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Derived")

    headers_s2 = [
        None, None, None,
        "Operating concern\n($CEERKRS)",
        "Controlling area\n($KOKRS)",
        "Company Code ($BUKRS)",
        "Plant / Valuation Area",
        "Profit Center",
        "Cost center",
        "Order type",
        "Asset Class",
        "Segment",
        "Business Place",
        "Credit segment",
        "Costing variant (K_KEKO)"
    ]

    for col_idx, h in enumerate(headers_s2, start=1):
        if h:
            c = ws2.cell(1, col_idx, value=h)
            c.font = header_font
            c.fill = header_fill
            c.alignment = center_align
            c.border = thin_border

    derived_rows = [
        ("Z_FI_PERIOD_CONTROL", "Z_FI_PERIOD_CONTROL_1000", "GL Period Control - CC 1000",
         "1000", "1000", "1000", "1000", "1000*", "1000*", "*", "*", "*", "*", "*", "*"),
        ("Z_FI_PERIOD_CONTROL", "Z_FI_PERIOD_CONTROL_2000", "GL Period Control - CC 2000",
         "1000", "1000", "2000", "2000", "2000*", "2000*", "*", "*", "*", "*", "*", "*"),
        ("Z_FI_GL_MASTER_MAINT", "Z_FI_GL_MASTER_MAINT_1000", "GL Master Data Maint - CC 1000",
         "1000", "1000", "1000", "1000", "1000*", "1000*", "*", "*", "*", "*", "*", "*"),
        ("Z_FI_GL_MASTER_MAINT", "Z_FI_GL_MASTER_MAINT_2000", "GL Master Data Maint - CC 2000",
         "1000", "1000", "2000", "2000", "2000*", "2000*", "*", "*", "*", "*", "*", "*"),
        ("Z_FI_JOURNAL_POST", "Z_FI_JOURNAL_POST_1000", "General Ledger Posting - CC 1000",
         "1000", "1000", "1000", "1000", "1000*", "1000*", "*", "*", "*", "*", "*", "*"),
        ("Z_FI_JOURNAL_POST", "Z_FI_JOURNAL_POST_2000", "General Ledger Posting - CC 2000",
         "1000", "1000", "2000", "2000", "2000*", "2000*", "*", "*", "*", "*", "*", "*"),
    ]

    for row_idx, r_data in enumerate(derived_rows, start=2):
        for col_idx, val in enumerate(r_data, start=1):
            c = ws2.cell(row_idx, col_idx, value=val)
            if col_idx in (1, 2):
                c.font = role_name_font
                c.fill = role_name_fill
            else:
                c.font = data_font
                if col_idx > 3:
                    c.alignment = center_align
            c.border = thin_border

    ws2.column_dimensions['A'].width = 28
    ws2.column_dimensions['B'].width = 32
    ws2.column_dimensions['C'].width = 35
    for col_idx in range(4, 16):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws2.column_dimensions[col_letter].width = 24

    # -------------------------------------------------------------
    # Sheet 3: Composite Roles
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Composite")

    comp_roles = [
        {"name": "Z_FI_COMP_CFO_1000", "desc": "Chief Financial Officer - CC 1000", "singles": ["Z_FI_PERIOD_CONTROL_1000", "Z_FI_JOURNAL_POST_1000"]},
        {"name": "Z_FI_COMP_ACCOUNTANT_1000", "desc": "GL Accountant - CC 1000", "singles": ["Z_FI_JOURNAL_POST_1000", "Z_FI_GL_MASTER_MAINT_1000"]},
        {"name": "Z_FI_COMP_CONTROLLER_1000", "desc": "Financial Controller - CC 1000", "singles": ["Z_FI_PERIOD_CONTROL_1000", "Z_FI_GL_MASTER_MAINT_1000"]},
        {"name": "Z_FI_COMP_EMPTY", "desc": "Empty Composite Role", "singles": []}
    ]

    all_singles = sorted(list({s for cr in comp_roles for s in cr["singles"]}))

    c_hdr3 = ws3.cell(1, 1, value="Single Role / Composite")
    c_hdr3.font = header_font
    c_hdr3.fill = header_fill
    c_hdr3.border = thin_border

    for col_idx, cr in enumerate(comp_roles, start=2):
        c_desc = ws3.cell(1, col_idx, value=cr["desc"])
        c_desc.font = header_font
        c_desc.fill = header_fill
        c_desc.border = thin_border

        c_name = ws3.cell(2, col_idx, value=cr["name"])
        c_name.font = role_name_font
        c_name.fill = role_name_fill
        c_name.border = thin_border

    for row_idx, s in enumerate(all_singles, start=3):
        c_s = ws3.cell(row_idx, 1, value=s)
        c_s.font = role_name_font
        c_s.fill = role_name_fill
        c_s.border = thin_border

        for col_idx, cr in enumerate(comp_roles, start=2):
            val = "X" if s in cr["singles"] else ""
            c_val = ws3.cell(row_idx, col_idx, value=val)
            c_val.font = data_font
            c_val.alignment = center_align
            c_val.border = thin_border

    ws3.column_dimensions['A'].width = 32
    for col_letter in ['B', 'C', 'D', 'E']:
        ws3.column_dimensions[col_letter].width = 32

    xlsx_path = os.path.join(OUT_DIR, "Template_Role.xlsx")
    wb.save(xlsx_path)
    print(f"[OK] Generated consolidated role template: {xlsx_path}")

    # Auto-convert to all converted files
    try:
        from convert_role_01_master import convert_role_01_master
        from convert_role_02_derive import convert_role_02_derive
        from convert_role_03_org_obj import convert_role_03_org_obj
        from convert_role_04_composite import convert_role_04_composite
        convert_role_01_master(xlsx_path)
        convert_role_02_derive(xlsx_path)
        convert_role_03_org_obj(xlsx_path)
        convert_role_04_composite(xlsx_path)
    except Exception as e:
        print(f"[WARN] Could not auto-convert Template_Role.xlsx: {e}")


# --- 2. Users Creation Template (Template_User_01_Create.xlsx & CSV) ---
def create_users_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "User Asignment"

    headers = ["FULL NAME", "Email", "Phone No.", "Department", "Function"]
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(1, col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = left_align

    sample_users = [
        ["John Doe", "johnd@company.com", "+84 981 234 567", "Finance & Accounting", "Finance Controller"],
        ["Jane Smith", "janes@company.com", "+84 971 888 999", "Finance & Accounting", "Senior Accountant"],
        ["Alex Johnson", "alexj@company.com", "+996 555 123 456", "PMU", "Project Director"],
        ["Michael Brown", "michaelb@company.com", "+996 500 654 321", "PMU", "Civil Engineer"],
        ["Emily Davis", "emilyd", None, "Administration", "HR Specialist"],
        ["David Wilson", "davidw@company.com", "0987654321", "operation", "OM"],
    ]

    for row_idx, r in enumerate(sample_users, start=2):
        for col_idx, val in enumerate(r, start=1):
            c = ws.cell(row_idx, col_idx, value=val)
            c.font = data_font
            c.border = thin_border

    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 22
    ws.column_dimensions['D'].width = 25
    ws.column_dimensions['E'].width = 30

    xlsx_path = os.path.join(OUT_DIR, "Template_User_01_Create.xlsx")
    wb.save(xlsx_path)
    print(f"[OK] Generated: {xlsx_path}")

    try:
        from convert_user_01_create import convert_user_01_create
        convert_user_01_create(xlsx_path)
    except ImportError:
        import subprocess
        script_path = os.path.join(os.path.dirname(__file__), "convert_user_01_create.py")
        subprocess.run([sys.executable, script_path, xlsx_path], check=True)


# --- 3. User Role Assignment Template (Template_User_02_Assign.xlsx & CSV) ---
def create_user_assign_role_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "User Asignment"

    roles = [
        "CR_COMMON",
        "Z_FI_COMP_CFO_1000",
        "Z_FI_COMP_ACCOUNTANT_1000",
        "Z_FI_COMP_CONTROLLER_1000",
        "Z_MM_PURCHASER_1000"
    ]

    ws.cell(1, 1, value="User ID").font = header_font
    ws.cell(1, 1).fill = header_fill
    ws.cell(1, 1).alignment = left_align

    for col_idx, rname in enumerate(roles, start=2):
        c = ws.cell(1, col_idx, value=rname)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center_align

    user_matrix = [
        ("JOHND", ["CR_COMMON", "Z_FI_COMP_CFO_1000"]),
        ("JANES", ["CR_COMMON", "Z_FI_COMP_ACCOUNTANT_1000"]),
        ("ALEXJ", ["CR_COMMON", "Z_FI_COMP_CONTROLLER_1000"]),
        ("MICHAELB", ["CR_COMMON", "Z_MM_PURCHASER_1000"]),
        ("EMILYD", ["CR_COMMON"]),
        ("DAVIDW", ["CR_COMMON", "Z_MM_PURCHASER_1000"]),
        ("EMPTY_USER", []),
    ]

    for row_idx, (uname, assigned_roles) in enumerate(user_matrix, start=2):
        c_user = ws.cell(row_idx, 1, value=uname)
        c_user.font = role_name_font
        c_user.fill = role_name_fill
        c_user.border = thin_border

        for col_idx, rname in enumerate(roles, start=2):
            val = "X" if rname in assigned_roles else ""
            c_val = ws.cell(row_idx, col_idx, value=val)
            c_val.font = data_font
            c_val.alignment = center_align
            c_val.border = thin_border

    ws.column_dimensions['A'].width = 22
    for col_letter in ['B', 'C', 'D', 'E', 'F']:
        ws.column_dimensions[col_letter].width = 30

    xlsx_path = os.path.join(OUT_DIR, "Template_User_02_Assign.xlsx")
    wb.save(xlsx_path)
    print(f"[OK] Generated: {xlsx_path}")

    try:
        from convert_user_02_assign import convert_user_02_assign
        convert_user_02_assign(xlsx_path)
    except ImportError:
        import subprocess
        script_path = os.path.join(os.path.dirname(__file__), "convert_user_02_assign.py")
        subprocess.run([sys.executable, script_path, xlsx_path], check=True)


# --- CLI Entrypoint ---
if __name__ == "__main__":
    print("\n--- Generating SAP Role & User Templates ---")
    create_consolidated_role_template()
    create_users_template()
    create_user_assign_role_template()
    print("\n[OK] All consolidated role & user templates and sample converted CSVs generated successfully.")
